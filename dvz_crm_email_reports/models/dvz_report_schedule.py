# -*- coding: utf-8 -*-
import base64
import io
from datetime import timedelta

from odoo import api, fields, models
from odoo.exceptions import UserError

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - handled at call time with a
    # clear UserError instead, so a missing dependency doesn't break
    # module installation itself, only the actual export/send action.
    openpyxl = None


def _dvz_sales_column(lead):
    """"Sales" column = ONLY the record's actual Salesperson (user_id,
    res.users) - the same field the per-salesperson sheets/emails are
    grouped by. Deliberately does NOT include sales_ids (the separate
    hr.employee many2many field) - showing both was confusing when a
    lead has both set to different people."""
    return lead.user_id.name if lead.user_id else ""


def _dvz_report_columns():
    """(header, getter, width) triples, in order, matching the reference
    Excel tracking sheet's columns - EXCEPT "PO - Ref #", deliberately
    dropped: that field (po_ref on crm.lead, from dvz_crm_ext) is used
    to store "Quote" instead of a distinct PO reference, so it isn't a
    separate report column here (it's what "Quote" below actually is).
    A function (not a module-level constant) so it's evaluated fresh
    each time, avoiding any stale-lambda-closure surprises. Width is in
    Excel's own character-count units - Customer gets the widest column
    (it usually has the longest text: full company names), the rest get
    a size that suits their typical content."""
    return [
        ("Quote", lambda lead: lead.po_ref or "", 14),
        ("Date", lambda lead: lead.inquiry_date or "", 12),
        ("Customer", lambda lead: lead.partner_id.name or "", 55),
        ("Project", lambda lead: ", ".join(lead.dvz_project.mapped("name")), 28),
        ("System", lambda lead: ", ".join(lead.system_ids.mapped("name")), 16),
        ("Brand", lambda lead: ", ".join(lead.brand_ids.mapped("name")), 16),
        ("Area", lambda lead: ", ".join(lead.area_ids.mapped("name")), 14),
        ("Sales", _dvz_sales_column, 24),
        ("Pre-Sales", lambda lead: lead.presales_id.name or "", 18),
        ("Value (w/o VAT)", lambda lead: lead.expected_revenue or 0.0, 16),
        ("Status", lambda lead: lead.stage_id.name or "", 14),
        ("Remarks", lambda lead: lead.remarks or "", 32),
        ("Quotation Deadline", lambda lead: lead.due_date or "", 18),
        ("Reason For Lost", lambda lead: lead.lost_reason_id.name or "", 26),
    ]


class DvzCrmReportSchedule(models.Model):
    _name = "dvz.crm.report.schedule"
    _description = "CRM Excel Report - Schedule / Email Configuration"

    name = fields.Char(required=True, default="Aala Tech Leads Report")
    active = fields.Boolean(default=True)

    # --- Who this schedule covers (which LEADS get included) ---
    team_ids = fields.Many2many(
        "crm.team", string="Teams",
        help="If set, every lead belonging to any of these Teams is "
             "included. Can be combined with Salespersons below - a "
             "lead matching EITHER counts.",
    )
    salesperson_ids = fields.Many2many(
        "res.users", "dvz_report_schedule_salesperson_rel",
        string="Salespersons",
        help="Only relevant when 'Send to Head' is OFF: narrows down "
             "WHICH salespersons get their own individual email. Leave "
             "empty in that mode to email EVERY salesperson found "
             "among the matched leads, each their own leads only. "
             "Also still used as an extra lead-matching filter either "
             "way (a lead matching Teams OR Salespersons is included).",
    )

    date_range_mode = fields.Selection(
        [("all", "All Leads"), ("range", "Specific Date Range")],
        string="Leads to Include", default="all", required=True,
    )
    date_from = fields.Date(string="From")
    date_to = fields.Date(string="To")

    # --- Who receives the COMBINED (per-salesperson-sheet) report ---
    send_to_head = fields.Boolean(
        string="Send to Head",
        help="ON: one combined workbook (one sheet per salesperson) "
             "goes to Head Recipients only - individual salespersons "
             "do NOT get a separate email in this mode. "
             "OFF: each salesperson found among the matched leads gets "
             "their own individual email with only their own leads "
             "(narrow this down with the Salespersons field above).",
    )
    head_user_ids = fields.Many2many(
        "res.users", "dvz_report_schedule_head_rel", string="Head Recipients",
        help="Receives the combined per-salesperson-sheet workbook - "
             "only used if 'Send to Head' is checked.",
    )

    # --- Automatic sending ---
    frequency_per_day = fields.Integer(
        string="Times per Day", default=1,
        help="How many times per day this schedule sends automatically. "
             "E.g. 2 sends roughly every 12 hours. The shared cron "
             "checks every 30 minutes and only actually sends once "
             "enough time has passed since this schedule's last send.",
    )
    last_sent_at = fields.Datetime(string="Last Sent", readonly=True)

    # ------------------------------------------------------------------
    # Lead selection
    # ------------------------------------------------------------------

    def _dvz_get_lead_domain(self):
        self.ensure_one()
        domain = []
        if self.date_range_mode == "range":
            if self.date_from:
                domain.append(("inquiry_date", ">=", self.date_from))
            if self.date_to:
                domain.append(("inquiry_date", "<=", self.date_to))

        team_or_person = []
        if self.team_ids:
            team_or_person.append(("team_id", "in", self.team_ids.ids))
        if self.salesperson_ids:
            team_or_person.append(("user_id", "in", self.salesperson_ids.ids))

        if len(team_or_person) == 2:
            domain = domain + ["|", team_or_person[0], team_or_person[1]]
        elif team_or_person:
            domain = domain + team_or_person
        # If NEITHER Teams nor Salespersons are set, no team/person
        # filter is applied at all - only the date range (if any)
        # restricts things, i.e. it behaves like "everyone".
        return domain

    def _dvz_get_leads(self):
        self.ensure_one()
        return self.env["crm.lead"].search(self._dvz_get_lead_domain())

    # ------------------------------------------------------------------
    # Excel building
    # ------------------------------------------------------------------

    def _dvz_check_openpyxl(self):
        if openpyxl is None:
            raise UserError(
                "The 'openpyxl' Python package isn't installed on this "
                "server. Ask your admin to run: "
                "pip install openpyxl --break-system-packages"
            )

    def _dvz_write_sheet(self, ws, leads):
        columns = _dvz_report_columns()
        header_font = Font(bold=True)
        header_fill = PatternFill("solid", fgColor="DDEBF7")
        thin = Side(style="thin", color="999999")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for col_idx, (header, _getter, _width) in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border

        for row_idx, lead in enumerate(leads, start=2):
            for col_idx, (_header, getter, _width) in enumerate(columns, start=1):
                try:
                    value = getter(lead)
                except Exception:
                    value = ""
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = border

        for col_idx, (_header, _getter, width) in enumerate(columns, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width

    def _dvz_build_excel(self, leads, sheet_title="Aala Tech"):
        """One plain workbook, one sheet, for the given leads."""
        self._dvz_check_openpyxl()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_title
        self._dvz_write_sheet(ws, leads)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def _dvz_build_excel_per_salesperson(self, leads):
        """One workbook, ONE SHEET PER SALESPERSON (by each lead's own
        Salesperson/user_id), each containing only that person's leads -
        used for the "Send to Head" email."""
        self._dvz_check_openpyxl()
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # drop the default blank sheet

        by_person = {}
        for lead in leads:
            person_name = lead.user_id.name if lead.user_id else "Unassigned"
            by_person.setdefault(person_name, self.env["crm.lead"])
            by_person[person_name] |= lead

        for person_name in sorted(by_person.keys()):
            # Excel sheet-name rules: max 31 chars, no [ ] : * ? / \
            safe_name = "".join(
                c for c in person_name if c not in "[]:*?/\\"
            )[:31] or "Sheet"
            ws = wb.create_sheet(title=safe_name)
            self._dvz_write_sheet(ws, by_person[person_name])

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # ------------------------------------------------------------------
    # Manual actions
    # ------------------------------------------------------------------

    def action_download_excel(self):
        """On-demand download - generates the current matching leads
        right now, independent of email/frequency settings entirely."""
        self.ensure_one()
        leads = self._dvz_get_leads()
        content = self._dvz_build_excel(leads)
        attachment = self.env["ir.attachment"].create({
            "name": f"Aala Tech - {self.name or 'CRM'} Leads.xlsx",
            "type": "binary",
            "datas": base64.b64encode(content),
            "res_model": self._name,
            "res_id": self.id,
        })
        return {
            "type": "ir.actions.act_url",
            "url": f"/web/content/{attachment.id}?download=true",
            "target": "self",
        }

    def action_send_now(self):
        """Manually trigger an immediate send, bypassing the frequency
        check - useful to test a schedule before leaving it to the
        cron."""
        for schedule in self:
            schedule._dvz_send_report()

    # ------------------------------------------------------------------
    # Sending
    # ------------------------------------------------------------------

    def _dvz_send_report(self):
        self.ensure_one()
        leads = self._dvz_get_leads()
        if not leads:
            return

        Mail = self.env["mail.mail"]

        if self.send_to_head:
            # ONE combined workbook, one sheet per salesperson, sent to
            # the Head recipient(s). Individual salespersons do NOT also
            # get a separate email in this mode - it's one or the other.
            if not self.head_user_ids:
                return
            content = self._dvz_build_excel_per_salesperson(leads)
            attachment = self.env["ir.attachment"].create({
                "name": "Aala Tech - All Salespersons Leads.xlsx",
                "type": "binary",
                "datas": base64.b64encode(content),
                "res_model": self._name,
                "res_id": self.id,
            })
            emails = [u.email for u in self.head_user_ids if u.email]
            if emails:
                Mail.create({
                    "subject": "Aala Tech - All Salespersons Leads",
                    "email_to": ",".join(emails),
                    "body_html": (
                        f"<p>All salespersons' leads are here - "
                        f"attached, {len(leads)} lead(s) total, one "
                        f"sheet per salesperson.</p>"
                    ),
                    "attachment_ids": [(6, 0, [attachment.id])],
                }).send()

        else:
            # Individual mode: one email PER salesperson found among
            # the matched leads (not just an explicitly selected
            # subset - if Salespersons is set, it narrows down WHICH
            # of them get emailed; if left empty, EVERY salesperson
            # represented in the results gets their own email).
            by_person = {}
            for lead in leads:
                if not lead.user_id:
                    continue
                by_person.setdefault(lead.user_id, self.env["crm.lead"])
                by_person[lead.user_id] |= lead

            for user, own_leads in by_person.items():
                if self.salesperson_ids and user not in self.salesperson_ids:
                    continue
                if not user.email:
                    continue
                content = self._dvz_build_excel(own_leads, sheet_title=user.name[:31])
                attachment = self.env["ir.attachment"].create({
                    "name": f"Aala Tech - {user.name} Leads.xlsx",
                    "type": "binary",
                    "datas": base64.b64encode(content),
                    "res_model": self._name,
                    "res_id": self.id,
                })
                Mail.create({
                    "subject": f"Aala Tech - {user.name} Leads",
                    "email_to": user.email,
                    "body_html": (
                        f"<p>Hi {user.name},</p>"
                        f"<p>Your leads are here - attached "
                        f"({len(own_leads)} lead(s)).</p>"
                    ),
                    "attachment_ids": [(6, 0, [attachment.id])],
                }).send()

        self.last_sent_at = fields.Datetime.now()

    @api.model
    def _cron_check_and_send_reports(self):
        """Called by ir.cron on a short, fixed interval (every 30
        minutes) - checks EACH active schedule's own frequency_per_day
        and only actually sends once enough time has passed since that
        schedule's last send. This lets one shared cron serve schedules
        with different frequencies, instead of needing a separate cron
        job per schedule."""
        now = fields.Datetime.now()
        for schedule in self.search([("active", "=", True)]):
            if schedule.frequency_per_day <= 0:
                continue
            interval = timedelta(hours=24.0 / schedule.frequency_per_day)
            if not schedule.last_sent_at or (now - schedule.last_sent_at) >= interval:
                schedule._dvz_send_report()
